"""Fase 2 — parser de `.docx` com paginação aproximada (handoff §5.1).

O .docx não carrega paginação real; estimamos páginas por volume de texto
(~3000 caracteres/página) e mapeamos qualquer posição para a página estimada.
"""
from __future__ import annotations

import hashlib
import math
from dataclasses import dataclass
from pathlib import Path

from docx import Document

CHARS_POR_PAGINA = 3000


@dataclass(frozen=True)
class DocumentoAnalisado:
    texto: str
    n_paginas: int


def checksum_arquivo(caminho: str | Path) -> str:
    """SHA-256 do conteúdo — base da idempotência de versões."""
    h = hashlib.sha256()
    with Path(caminho).open("rb") as f:
        for bloco in iter(lambda: f.read(1 << 16), b""):
            h.update(bloco)
    return h.hexdigest()


def parse_docx(caminho: str | Path) -> DocumentoAnalisado:
    """Extrai texto de parágrafos e tabelas, com estimativa de páginas."""
    doc = Document(str(caminho))
    partes: list[str] = []
    for paragrafo in doc.paragraphs:
        texto = paragrafo.text.strip()
        if texto:
            partes.append(texto)
    for tabela in doc.tables:
        for linha in tabela.rows:
            for celula in linha.cells:
                texto = celula.text.strip()
                if texto:
                    partes.append(texto)
    texto = "\n".join(partes)
    n_paginas = max(1, math.ceil(len(texto) / CHARS_POR_PAGINA))
    return DocumentoAnalisado(texto=texto, n_paginas=n_paginas)


def pagina_estimada(texto: str, posicao: int, n_paginas: int) -> int:
    """Página estimada (1-based) de uma posição no texto."""
    if not texto:
        return 1
    return min(n_paginas, 1 + posicao // CHARS_POR_PAGINA)
