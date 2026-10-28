"""Testes sintéticos de regressão. Não substituem gabarito de currículos reais."""
import io
import unittest
from docx import Document
from src.experimentos.analisar_curriculos import detect_features, extract_docx

class TestarExtracao(unittest.TestCase):
    def test_alias_e_limiar(self):
        v=detect_features('ReactJS, Postgre SQL, JavaScript, C++ e AWS.\n')
        for term in ('React','PostgreSQL','JavaScript','C++','AWS'):
            self.assertIn(term,v['skills'])
        self.assertNotIn('Java',v['skills'])
    def test_java_isolado(self):
        self.assertIn('Java',detect_features('Core Java 17')['skills'])
    def test_nao_inferir_certificados(self):
        self.assertNotIn('PMP',detect_features('Project management with Scrum')['skills'])
    def test_leitura_de_tabela_e_cabecalho(self):
        doc=Document()
        doc.add_paragraph('TESTE DOCX')
        doc.sections[0].header.paragraphs[0].text='EXEMPLO HEADER'
        t=doc.add_table(rows=1,cols=2)
        t.cell(0,0).text='Competência'
        t.cell(0,1).text='Kubernetes'
        buf=io.BytesIO();doc.save(buf)
        txt,pars,tables,table_pars=extract_docx(buf.getvalue())
        self.assertIn('Kubernetes',txt)
        self.assertIn('EXEMPLO HEADER',txt)
        self.assertEqual(tables,1)
        self.assertGreaterEqual(table_pars,2)

if __name__=='__main__':unittest.main()
