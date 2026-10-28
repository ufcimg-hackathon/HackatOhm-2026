"""Fase 2 — testes do pipeline de ingestão."""
from pathlib import Path

import pytest
from docx import Document

from src.extraction.parser_docx import parse_docx
from src.governance.auditoria import ultimos
from src.governance.vault import CofrePII, buscar_pii
from src.ingestion.db import conectar, criar_schema
from src.ingestion.pipeline import (
    carregar_inventario,
    ingerir_inventario,
    ingerir_registro,
)


@pytest.fixture()
def con(tmp_path):
    con = conectar(tmp_path / "pipeline.db")
    criar_schema(con)
    yield con
    con.close()


@pytest.fixture()
def cofre():
    return CofrePII(CofrePII.gerar_chave())


def _registro(arquivo: Path, indice: str = "1") -> dict:
    return {
        "id": indice,
        "nome_arquivo": arquivo.name,
        "nome_candidato": arquivo.stem,
        "link_individual": "",
        "status_download": "ok (cache)",
        "caminho_cache": str(arquivo).replace("\\", "/"),
    }


def _primeiro_docx_com_texto() -> Path:
    cache = Path("data/cache_drive")
    arquivos = sorted(cache.glob("*.docx")) if cache.exists() else []
    for arquivo in arquivos[:20]:
        if len(parse_docx(arquivo).texto) >= 200:
            return arquivo
    pytest.skip("cache de curriculos indisponivel neste ambiente")


def _docx_sintetico(diretorio: Path, nome: str, email: str, extra: str = "") -> Path:
    doc = Document()
    doc.add_paragraph(f"Curriculo de teste - contato {email}")
    doc.add_paragraph(
        "Experiencia: Analista de Sistemas (2019 - 2022) com Java, SQL e Spring."
    )
    doc.add_paragraph("Formacao: Bacharelado em Computacao - Universidade Teste.")
    if extra:
        doc.add_paragraph(extra)
    caminho = diretorio / nome
    doc.save(caminho)
    return caminho


def test_ingere_um_docx_real_e_redige_pii(con, cofre):
    arquivo = _primeiro_docx_com_texto()
    original = parse_docx(arquivo).texto
    resultado = ingerir_registro(con, cofre, _registro(arquivo))
    assert resultado["status"] == "processado"

    curriculo = con.execute("SELECT * FROM curriculos").fetchone()
    assert curriculo["status_processamento"] == "processado"
    assert curriculo["n_paginas"] >= 1
    assert curriculo["checksum"]
    raw = curriculo["raw_texto"]
    assert raw

    for achado in buscar_pii(original):
        valor = str(achado["valor"])
        if valor and len(valor) >= 5:
            assert valor not in raw, f"PII vazou no raw_texto: {achado['tipo']}"

    # PII que existia no original foi para o cofre cifrado
    if buscar_pii(original):
        assert con.execute(
            "SELECT COUNT(*) c FROM cofre_dados_pessoais"
        ).fetchone()["c"] >= 1


def test_ingestao_e_idempotente_por_checksum(con, cofre):
    arquivo = _primeiro_docx_com_texto()
    primeiro = ingerir_registro(con, cofre, _registro(arquivo))
    segundo = ingerir_registro(con, cofre, _registro(arquivo))
    assert primeiro["status"] == "processado"
    assert segundo["status"] == "duplicado"
    assert con.execute("SELECT COUNT(*) c FROM curriculos").fetchone()["c"] == 1
    assert con.execute("SELECT COUNT(*) c FROM candidatos").fetchone()["c"] == 1


def test_mesmo_email_gera_mesmo_candidato(con, cofre, tmp_path):
    email = "candidato.teste@exemplo.com"
    primeiro = _docx_sintetico(tmp_path, "a.docx", email, "Curso de Java avancado.")
    segundo = _docx_sintetico(
        tmp_path, "b.docx", email, "Certificacao PMP; projeto de energia."
    )
    mapa = {}
    r1 = ingerir_registro(con, cofre, _registro(primeiro, "1"), mapa_emails=mapa)
    r2 = ingerir_registro(con, cofre, _registro(segundo, "2"), mapa_emails=mapa)
    assert r1["status"] == "processado" and r2["status"] == "processado"
    assert r1["candidato_criado"] is True
    assert r2["candidato_criado"] is False
    assert con.execute("SELECT COUNT(*) c FROM candidatos").fetchone()["c"] == 1
    # segundo currículo do mesmo candidato = versão 2
    versoes = [
        row["versao"]
        for row in con.execute("SELECT versao FROM curriculos ORDER BY id")
    ]
    assert versoes == [1, 2]


def test_pii_do_docx_sintetico_vai_para_cofre_e_sai_do_texto(con, cofre, tmp_path):
    email = "pii.teste@exemplo.com"
    arquivo = _docx_sintetico(tmp_path, "pii.docx", email)
    ingerir_registro(con, cofre, _registro(arquivo))
    raw = con.execute("SELECT raw_texto FROM curriculos").fetchone()["raw_texto"]
    assert email not in raw
    assert "***" in raw
    linha = con.execute(
        "SELECT * FROM cofre_dados_pessoais WHERE tipo = 'email'"
    ).fetchone()
    assert linha is not None and cofre.decifrar(linha["valor_criptografado"]) == email


def test_lote_respeita_limite_e_audita(con, cofre):
    csv = Path("data/inventario_curriculos.csv")
    if not csv.exists() or not Path("data/cache_drive").exists():
        pytest.skip("inventario/cache indisponiveis neste ambiente")
    resumo = ingerir_inventario(con, cofre, csv, limite=3)
    assert resumo.total == 3
    assert resumo.processados + resumo.erros + resumo.duplicados == 3
    assert resumo.processados >= 1
    acoes = [row["acao"] for row in ultimos(con, limite=50)]
    assert "ingerir_curriculo" in acoes


def test_atualizacao_forcada_cria_nova_versao_do_mesmo_candidato(con, cofre, tmp_path):
    email = "atualiza.teste@exemplo.com"
    primeiro = _docx_sintetico(tmp_path, "v1.docx", email)
    r1 = ingerir_registro(con, cofre, _registro(primeiro, "1"))
    candidato_id = r1["candidato_id"]
    con.execute(
        "UPDATE candidatos SET atualizado_em = '2000-01-01 00:00:00' WHERE id = ?",
        (candidato_id,),
    )
    con.commit()

    segundo = _docx_sintetico(
        tmp_path, "v2.docx", email, "Projeto de energia na EDP com Java e SQL."
    )
    r2 = ingerir_registro(
        con, cofre, _registro(segundo, "2"), candidato_id_forcado=candidato_id
    )
    assert r2["status"] == "processado"
    assert r2["candidato_id"] == candidato_id
    assert r2["versao"] == 2
    assert con.execute("SELECT COUNT(*) c FROM candidatos").fetchone()["c"] == 1
    versoes = [
        row["versao"]
        for row in con.execute(
            "SELECT versao FROM curriculos WHERE candidato_id = ? ORDER BY id",
            (candidato_id,),
        )
    ]
    assert versoes == [1, 2]
    atualizado = con.execute(
        "SELECT atualizado_em FROM candidatos WHERE id = ?", (candidato_id,)
    ).fetchone()["atualizado_em"]
    assert atualizado != "2000-01-01 00:00:00"
    acoes = [row["acao"] for row in ultimos(con, limite=20)]
    assert "atualizar_curriculo" in acoes

    # o mesmo arquivo não cria versão 3
    r3 = ingerir_registro(
        con, cofre, _registro(segundo, "3"), candidato_id_forcado=candidato_id
    )
    assert r3["status"] == "duplicado"
    assert con.execute("SELECT COUNT(*) c FROM curriculos").fetchone()["c"] == 2


def test_atualizacao_recusa_email_de_outro_candidato(con, cofre, tmp_path):
    primeiro = _docx_sintetico(tmp_path, "a.docx", "a@exemplo.com")
    r1 = ingerir_registro(con, cofre, _registro(primeiro, "1"))
    segundo = _docx_sintetico(tmp_path, "b.docx", "b@exemplo.com")
    ingerir_registro(con, cofre, _registro(segundo, "2"))
    outro = _docx_sintetico(tmp_path, "c.docx", "b@exemplo.com")
    with pytest.raises(ValueError):
        ingerir_registro(
            con,
            cofre,
            _registro(outro, "3"),
            candidato_id_forcado=r1["candidato_id"],
        )


def test_carregar_inventario_so_ok(tmp_path):
    csv = tmp_path / "inv.csv"
    csv.write_text(
        "id,nome_arquivo,nome_candidato,link_individual,status_download,caminho_cache\n"
        "1,a.docx,A,,ok (cache),data/a.docx\n"
        "2,b.docx,B,,falhou,data/b.docx\n",
        encoding="utf-8",
    )
    registros = carregar_inventario(csv)
    assert [r["id"] for r in registros] == ["1"]
