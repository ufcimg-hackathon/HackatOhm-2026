"""Etapa 0 — integridade da base, extração, PII e alocação."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from scripts.build_db import GATES_ENERGIA_V1, PESOS_ENERGIA_V1, construir
from scripts.reset_demo import criar_checkpoint, reset
from scripts.seed_demo import seed
from src.extraction.regras import extrair
from src.governance.vault import CofrePII, buscar_pii
from src.ingestion import inbox, pipeline
from src.ingestion.db import TABELAS_ESPERADAS, conectar, criar_schema, tabelas
from src.ingestion.pipeline import ingerir_registro
from src.scoring.servico import publicar_nova_versao, recalcular_disputa, rollback_preset
from src.scoring.time import salvar_alocacao


def _docx(diretorio: Path, nome: str, email: str, extra: str = "") -> Path:
    doc = Document()
    doc.add_paragraph(f"Curriculo de teste - contato {email}")
    doc.add_paragraph(
        "Experiencia: Analista de Sistemas - EnergiaPaulista (2019 - 2022)"
    )
    doc.add_paragraph("Projeto de smart grid com Java, SQL e Spring.")
    doc.add_paragraph("Formacao: Bacharelado em Computacao - Universidade Teste.")
    if extra:
        doc.add_paragraph(extra)
    caminho = diretorio / nome
    doc.save(caminho)
    return caminho


def _registro(arquivo: Path) -> dict:
    return {
        "id": arquivo.stem,
        "nome_arquivo": arquivo.name,
        "nome_candidato": arquivo.stem,
        "link_individual": "",
        "status_download": "ok (teste)",
        "caminho_cache": str(arquivo).replace("\\", "/"),
    }


def _candidato(con, nome: str, checksum: str) -> tuple[int, int]:
    candidato_id = int(
        con.execute(
            "INSERT INTO candidatos (nome_pseudonimo) VALUES (?)", (nome,)
        ).lastrowid
    )
    curriculo_id = int(
        con.execute(
            "INSERT INTO curriculos (candidato_id, nome_arquivo, caminho_cache,"
            " checksum, n_paginas, versao)"
            " VALUES (?, ?, ?, ?, 1, 1)",
            (candidato_id, f"{nome}.docx", "x", checksum),
        ).lastrowid
    )
    return candidato_id, curriculo_id


def test_criar_schema_e_idempotente(tmp_path):
    con = conectar(tmp_path / "schema.db")
    try:
        criar_schema(con)
        criar_schema(con)
        assert tabelas(con) == TABELAS_ESPERADAS
    finally:
        con.close()


def test_ingestao_falha_nao_deixa_curriculo_parcial(tmp_path, monkeypatch):
    arquivo = _docx(tmp_path, "transacao.docx", "transacao@exemplo.com")
    con = conectar(tmp_path / "transacao.db")
    criar_schema(con)
    cofre = CofrePII(CofrePII.gerar_chave())
    original = pipeline._persistir_fato

    def falhar(*args, **kwargs):
        raise RuntimeError("falha simulada no meio da extração")

    monkeypatch.setattr(pipeline, "_persistir_fato", falhar)
    try:
        with pytest.raises(RuntimeError):
            ingerir_registro(con, cofre, _registro(arquivo))
        for tabela in (
            "candidatos",
            "curriculos",
            "experiencias",
            "competencias",
            "evidencias",
            "cofre_dados_pessoais",
            "registros_auditoria",
        ):
            total = con.execute(f"SELECT COUNT(*) c FROM {tabela}").fetchone()["c"]
            assert total == 0, f"{tabela} ficou parcialmente gravada"
    finally:
        monkeypatch.setattr(pipeline, "_persistir_fato", original)
        con.close()

    # A repetição da operação recupera o arquivo do zero.
    con = conectar(tmp_path / "transacao.db")
    try:
        resultado = ingerir_registro(con, cofre, _registro(arquivo))
        assert resultado["status"] == "processado"
        assert con.execute("SELECT COUNT(*) c FROM curriculos").fetchone()["c"] == 1
    finally:
        con.close()


def test_pii_em_campos_estruturados_e_redigida(tmp_path):
    email = "rh.secreto@empresa.com"
    doc = Document()
    doc.add_paragraph(f"Curriculo de teste - contato {email}")
    doc.add_paragraph(
        f"Experiência: Analista de Sistemas - {email} (2019 - 2022)"
    )
    doc.add_paragraph("Projeto de energia com Java e SQL.")
    doc.add_paragraph("Formação: Bacharelado em Computação - Universidade Teste.")
    arquivo = tmp_path / "pii-estruturada.docx"
    doc.save(arquivo)
    con = conectar(tmp_path / "pii.db")
    criar_schema(con)
    cofre = CofrePII(CofrePII.gerar_chave())
    try:
        ingerir_registro(con, cofre, _registro(arquivo))
        linha = con.execute("SELECT * FROM experiencias").fetchone()
        assert linha is not None
        texto = " ".join(
            str(linha[campo] or "")
            for campo in ("empresa", "papel", "resumo", "tecnologias")
        )
        assert email not in texto
        assert "***" in texto, "o campo estruturado deveria estar redigido"
        for ev in con.execute("SELECT trecho FROM evidencias"):
            assert email not in (ev["trecho"] or "")
        for res in con.execute("SELECT trecho FROM residual"):
            assert email not in (res["trecho"] or "")
        raw = con.execute("SELECT raw_texto FROM curriculos").fetchone()["raw_texto"]
        assert email not in raw
    finally:
        con.close()


def test_telefone_nao_confunde_intervalo_de_anos():
    assert not [a for a in buscar_pii("Experiência 2015 - 2023") if a["tipo"] == "telefone"]
    achados = [a for a in buscar_pii("Contato (11) 98765-4321") if a["tipo"] == "telefone"]
    assert achados


def test_windows_nao_e_classificado_como_energia():
    ext = extrair(
        "Desenvolvedor Windows\n2019 - 2022\nSistemas Windows e .NET.",
        n_paginas=1,
    )
    assert len(ext.experiencias) == 1
    assert ext.experiencias[0].dados["dominio"] is None


def test_formacao_com_periodo_em_linha_separada_nao_vira_experiencia():
    texto = (
        "Formação\n"
        "Bacharelado em Ciência da Computação\n"
        "2010 - 2014\n"
        "\n"
        "Experiência\n"
        "Analista de Sistemas\n"
        "2015 - 2020\n"
        "Java e SQL."
    )
    ext = extrair(texto, n_paginas=1)
    assert len(ext.experiencias) == 1
    assert "2015" in ext.experiencias[0].dados["periodo"]
    assert all("2010" not in e.dados["periodo"] for e in ext.experiencias)


def test_alocacao_recusa_candidato_eliminado(tmp_path):
    db = tmp_path / "alocacao.db"
    construir(db)
    con = conectar(db)
    try:
        papeis = {
            row["sigla"]: int(row["id"])
            for row in con.execute("SELECT id, sigla FROM papeis")
        }
        candidato_id, curriculo_id = _candidato(con, "Eliminado", "eliminado")
        con.execute(
            "INSERT INTO disputas (nome_ficticio, preset_versao_id) VALUES ('Demo', 1)"
        )
        disputa_id = int(con.execute("SELECT id FROM disputas").fetchone()["id"])
        con.execute(
            "INSERT INTO pontuacoes (disputa_id, candidato_id, papel_id,"
            " versao_preset_id, curriculo_versao_id, pontuacao_ponderada,"
            " confianca, pontuacao_final, detalhe_json)"
            " VALUES (?, ?, ?, 1, ?, 10, 10, 10, ?)",
            (
                disputa_id,
                candidato_id,
                papeis["pm_scrum"],
                curriculo_id,
                json.dumps({"status_gates": "eliminado"}),
            ),
        )
        con.commit()
        with pytest.raises(ValueError, match="eliminado"):
            salvar_alocacao(
                con,
                disputa_id,
                [{"papel_id": papeis["pm_scrum"], "candidato_id": candidato_id}],
            )
        assert con.execute("SELECT COUNT(*) c FROM alocacoes").fetchone()["c"] == 0
    finally:
        con.close()


def test_rollback_nao_esconde_candidato_novo_sem_score_na_versao_alvo(tmp_path):
    db = tmp_path / "rollback.db"
    construir(db)
    con = conectar(db)
    try:
        primeiro_id, _ = _candidato(con, "Primeiro", "primeiro")
        con.execute(
            "INSERT INTO disputas (nome_ficticio, preset_versao_id) VALUES ('Demo', 1)"
        )
        disputa_id = int(con.execute("SELECT id FROM disputas").fetchone()["id"])
        recalcular_disputa(con, disputa_id, 1)

        segundo_id, _ = _candidato(con, "Segundo", "segundo")
        publicar_nova_versao(
            con,
            disputa_id,
            1,
            PESOS_ENERGIA_V1,
            GATES_ENERGIA_V1,
            ator="teste",
        )
        resultado = rollback_preset(con, disputa_id, 1, ator="teste")
        assert resultado["recalculou"] is True
        for candidato_id in (primeiro_id, segundo_id):
            total = con.execute(
                "SELECT COUNT(*) c FROM pontuacoes"
                " WHERE disputa_id = ? AND versao_preset_id = 1"
                " AND candidato_id = ?",
                (disputa_id, candidato_id),
            ).fetchone()["c"]
            assert total > 0
    finally:
        con.close()


def test_inbox_processa_duplica_e_permite_retry(tmp_path, monkeypatch):
    pasta = tmp_path / "inbox"
    pasta.mkdir()
    _docx(pasta, "a.docx", "a@exemplo.com")
    con = conectar(tmp_path / "inbox.db")
    criar_schema(con)
    cofre = CofrePII(CofrePII.gerar_chave())
    try:
        scan = inbox.escanear(con, pasta)
        assert scan[0]["status"] == "pendente"

        resumo = inbox.sincronizar(con, cofre, pasta)
        assert resumo["processados"] == 1
        assert resumo["erros"] == 0
        assert con.execute(
            "SELECT status FROM inbox_arquivos"
        ).fetchone()["status"] == "processado"

        resumo2 = inbox.sincronizar(con, cofre, pasta)
        assert resumo2["duplicados"] == 1
        assert con.execute("SELECT COUNT(*) c FROM curriculos").fetchone()["c"] == 1

        _docx(pasta, "b.docx", "b@exemplo.com")
        original = inbox.ingerir_registro

        def falhar(*args, **kwargs):
            raise RuntimeError("parse simulado")

        monkeypatch.setattr(inbox, "ingerir_registro", falhar)
        resumo_erro = inbox.sincronizar(con, cofre, pasta)
        assert resumo_erro["erros"] == 1
        linha_erro = con.execute(
            "SELECT status, erro FROM inbox_arquivos WHERE nome_arquivo = 'b.docx'"
        ).fetchone()
        assert linha_erro["status"] == "erro"
        assert "parse simulado" in linha_erro["erro"]

        monkeypatch.setattr(inbox, "ingerir_registro", original)
        resumo_retry = inbox.sincronizar(con, cofre, pasta)
        assert resumo_retry["processados"] == 1
        assert con.execute(
            "SELECT status FROM inbox_arquivos WHERE nome_arquivo = 'b.docx'"
        ).fetchone()["status"] == "processado"
    finally:
        con.close()


def test_seed_demo_aplica_curadoria_sem_inventar_lista(tmp_path):
    db = tmp_path / "seed.db"
    construir(db)
    con = conectar(db)
    try:
        primeiro, _ = _candidato(con, "Curado", "curado")
        segundo, _ = _candidato(con, "Nao Curado", "nao-curado")
        con.commit()
    finally:
        con.close()
    csv = tmp_path / "curadoria.csv"
    csv.write_text(
        "candidato_id\n" f"{primeiro}\n" "999999\n",
        encoding="utf-8",
    )
    resultado = seed(db, csv)
    assert resultado["marcados"] == 1
    assert len(resultado["ausentes"]) == 1
    con = conectar(db)
    try:
        status = {
            row["id"]: row["status_curacao"]
            for row in con.execute("SELECT id, status_curacao FROM candidatos")
        }
        assert status[primeiro] == "curado"
        assert status[segundo] == "nao_curado"
        assert con.execute(
            "SELECT COUNT(*) c FROM registros_auditoria WHERE acao = 'seed_curadoria'"
        ).fetchone()["c"] == 1
    finally:
        con.close()


def test_reset_demo_restaura_checkpoint(tmp_path):
    db = tmp_path / "demo.db"
    checkpoint = tmp_path / "demo_checkpoint.db"
    construir(db)
    criar_checkpoint(db, checkpoint)
    con = conectar(db)
    try:
        con.execute(
            "INSERT INTO candidatos (nome_pseudonimo) VALUES ('Depois do checkpoint')"
        )
        con.commit()
        assert con.execute("SELECT COUNT(*) c FROM candidatos").fetchone()["c"] == 1
    finally:
        con.close()
    reset(db, checkpoint, force=True)
    con = conectar(db)
    try:
        assert con.execute("SELECT COUNT(*) c FROM candidatos").fetchone()["c"] == 0
    finally:
        con.close()
