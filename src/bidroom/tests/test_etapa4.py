"""Etapa 4 — edital versionado, rastreável e revisado antes da ativação."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from src.extraction.edital import TIPOS, extrair_edital
from src.governance import editais
from src.governance.ciclo import listar_atualizacoes_pendentes
from src.ingestion.db import conectar, criar_schema

LINHAS_EDITAL = [
    "EDITAL DE CONTRATAÇÃO — TI ENERGIA",
    "2. Papéis e vagas",
    "Papel: PM/Scrum Sênior — 1 vaga",
    "Papel: BA/BSA Sênior — 1 vaga",
    "Papel: Java Developer Sênior — 2 vagas",
    "Papel: BI/Data Analyst Sênior — 1 vaga",
    "3. Requisitos individuais",
    "Competências obrigatórias: Gestão de projetos, Scrum, Kanban, Excel, Power BI",
    "Experiência mínima: 5 anos em gestão de projetos",
    "Certificação obrigatória: PMP",
    "Certificações desejáveis: CSM, PSM",
    "Comprovação: certificado PMP válido",
    "Período de trabalho: 12 meses",
    "Critério eliminatório: não ter disponibilidade para viagens",
    "Pontuação: 30 pontos para experiência em energia",
    "Limite: no máximo 2 candidatos por papel",
    "Regra de combinação: a equipe deve ter pelo menos um PMP e um CBAP",
    "4. Requisitos da empresa",
    "A empresa deve possuir certificação ISO 9001.",
    "5. Requisitos de equipe",
    "A equipe deve ter pelo menos 2 profissionais certificados PMP.",
    "6. Disposições",
    "Obrigatório: disponibilidade para viagens.",
    "Critério de pontuação: qualidade da proposta (a definir).",
]


def _docx(
    tmp_path: Path,
    linhas: list[str],
    *,
    nome: str = "edital.docx",
    tabela: list[list[str]] | None = None,
) -> Path:
    caminho = tmp_path / nome
    doc = Document()
    for linha in linhas:
        doc.add_paragraph(linha)
    if tabela:
        table = doc.add_table(rows=len(tabela), cols=len(tabela[0]))
        for r, linha in enumerate(tabela):
            for c, valor in enumerate(linha):
                table.rows[r].cells[c].text = valor
    doc.save(caminho)
    return caminho


@pytest.fixture()
def con(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("BIDROOM_EDITAIS", str(tmp_path / "editais"))
    conexao = conectar(":memory:")
    criar_schema(conexao)
    yield conexao
    conexao.close()


def _criar_e_receber(
    con,
    tmp_path: Path,
    linhas: list[str] | None = None,
    *,
    tabela: list[list[str]] | None = None,
) -> tuple[int, dict]:
    edital_id = editais.criar_edital(
        con, "Edital de teste", cliente_ficticio="Cliente Fictício"
    )
    caminho = _docx(tmp_path, linhas or LINHAS_EDITAL, tabela=tabela)
    resultado = editais.receber_versao(
        con,
        edital_id,
        nome_arquivo=caminho.name,
        conteudo=caminho.read_bytes(),
        ator="teste",
    )
    return edital_id, resultado


def _revisar_todos(con, versao_id: int, *, decisao: str = "aprovar") -> None:
    for requisito in editais.listar_requisitos(con, versao_id):
        editais.revisar_requisito(
            con,
            int(requisito["id"]),
            revisor="revisor",
            decisao=decisao,
            observacao="revisão de teste",
        )


def _receber_e_ativar(con, tmp_path: Path) -> tuple[int, int]:
    edital_id, recebido = _criar_e_receber(con, tmp_path)
    _revisar_todos(con, recebido["versao_id"])
    editais.aprovar_versao(con, recebido["versao_id"], revisor="revisor")
    editais.ativar_versao(con, recebido["versao_id"], ator="teste")
    return edital_id, int(recebido["versao_id"])


# ---------------------------------------------------------------- parser


def test_parser_compara_com_leitura_humana_e_preserva_rastreabilidade(tmp_path):
    caminho = _docx(
        tmp_path,
        LINHAS_EDITAL,
        tabela=[
            ["Papel", "Vagas", "Competências"],
            ["PM/Scrum Sênior", "1", "Scrum, Kanban"],
        ],
    )
    extraido = extrair_edital(caminho)
    requisitos = extraido.requisitos

    papeis = {
        req.regra["sigla"]: req.regra["vagas"]
        for req in requisitos
        if req.tipo == "papel"
    }
    assert papeis == {
        "pm_scrum": 1,
        "ba_bsa": 1,
        "java": 2,
        "bi_data": 1,
    }

    competencias = [req for req in requisitos if req.tipo == "competencia"]
    assert any(
        {"Scrum", "Kanban"}.issubset(set(req.regra["itens"]))
        for req in competencias
    )

    experiencia = next(
        req for req in requisitos if req.tipo == "experiencia_minima"
    )
    assert experiencia.regra["anos"] == 5
    assert experiencia.obrigatorio is True

    certificacoes = [req for req in requisitos if req.tipo == "certificacao"]
    pmp = next(req for req in certificacoes if "PMP" in req.regra["itens"])
    assert pmp.obrigatorio is True
    assert any(
        not req.obrigatorio and {"CSM", "PSM"}.issubset(set(req.regra["itens"]))
        for req in certificacoes
    )

    comprovante = next(req for req in requisitos if req.tipo == "comprovante")
    assert comprovante.obrigatorio is True
    assert "PMP" in comprovante.regra["itens"]

    periodo = next(req for req in requisitos if req.tipo == "periodo_trabalho")
    assert periodo.regra["duracao_meses"] == 12

    assert any(req.tipo == "criterio_eliminatorio" for req in requisitos)
    pontuacao = next(req for req in requisitos if req.tipo == "pontuacao")
    assert pontuacao.regra["pontos"] == 30
    limite = next(req for req in requisitos if req.tipo == "limite")
    assert limite.regra["valor"] == 2
    combinacao = next(req for req in requisitos if req.tipo == "regra_combinacao")
    assert combinacao.escopo == "equipe"

    empresa = [req for req in requisitos if req.escopo == "empresa"]
    assert empresa
    assert all(req.regra.get("pontua") is False for req in empresa)
    equipe = [req for req in requisitos if req.escopo == "equipe"]
    assert equipe
    assert all(req.regra.get("pontua") is False for req in equipe)

    # Nenhum requisito obrigatório pode desaparecer sem tratamento explícito.
    assert any(
        req.tipo == "outro"
        and req.obrigatorio
        and "viagens" in req.descricao
        for req in requisitos
    )
    assert any(req.ambiguidade for req in requisitos)

    for requisito in requisitos:
        assert requisito.tipo in TIPOS
        assert requisito.trecho.strip()
        assert requisito.localizacao.strip()
        assert requisito.pagina >= 1


def test_requisito_obrigatorio_desconhecido_nao_e_omitido(tmp_path):
    caminho = _docx(
        tmp_path,
        [
            "Obrigatório: disponibilidade para viagens.",
            "Desejável: inglês avançado.",
        ],
    )
    extraido = extrair_edital(caminho)
    outros = [req for req in extraido.requisitos if req.tipo == "outro"]
    assert len(outros) == 2
    obrigatorio = next(req for req in outros if "viagens" in req.descricao)
    opcional = next(req for req in outros if "inglês" in req.descricao)
    assert obrigatorio.obrigatorio is True
    assert opcional.obrigatorio is False
    assert all(req.ambiguidade for req in outros)


def test_tabela_dificil_gera_tratamento_explicito(tmp_path):
    caminho = _docx(
        tmp_path,
        ["Tabela de condições"],
        tabela=[
            ["Item", "Observação"],
            ["Requisito X", "a definir"],
        ],
    )
    extraido = extrair_edital(caminho)
    tratamento = next(
        req for req in extraido.requisitos if req.tipo == "tratamento_documento"
    )
    assert tratamento.ambiguidade is True
    assert "tabela 1" in tratamento.localizacao
    assert extraido.alertas


def test_documento_sem_texto_gera_tratamento_explicito(tmp_path):
    caminho = _docx(tmp_path, [])
    extraido = extrair_edital(caminho)
    tratamento = next(
        req for req in extraido.requisitos if req.tipo == "tratamento_documento"
    )
    assert tratamento.ambiguidade is True
    assert "OCR" in tratamento.descricao


# ---------------------------------------------------------------- serviço


def test_receber_versao_e_idempotente_por_checksum(con, tmp_path):
    edital_id, recebido = _criar_e_receber(con, tmp_path)
    assert recebido["status"] == "extraida"
    assert recebido["requisitos"] >= 10
    caminho = tmp_path / "edital.docx"
    duplicado = editais.receber_versao(
        con,
        edital_id,
        nome_arquivo=caminho.name,
        conteudo=caminho.read_bytes(),
        ator="teste",
    )
    assert duplicado["duplicado"] is True
    assert duplicado["versao_id"] == recebido["versao_id"]
    assert len(editais.listar_versoes(con, edital_id)) == 1


def test_ativacao_bloqueada_ate_revisar_tudo_e_resolver_ambiguidade(con, tmp_path):
    edital_id, recebido = _criar_e_receber(con, tmp_path)
    versao_id = int(recebido["versao_id"])
    with pytest.raises(ValueError, match="pendente"):
        editais.aprovar_versao(con, versao_id, revisor="revisor")

    # Revisa todos; a aprovação explícita resolve a ambiguidade.
    _revisar_todos(con, versao_id)
    assert editais.status_revisao(con, versao_id)["ambiguos_pendentes"] == 0
    assert editais.aprovar_versao(con, versao_id, revisor="revisor")["status"] == "aprovada"
    assert editais.ativar_versao(con, versao_id, ator="teste")["status"] == "ativa"

    # Invariante independente do caminho feliz: ambiguidade marcada continua
    # bloqueando mesmo que alguém tente forçar o status por SQL.
    edital2 = editais.criar_edital(con, "Edital com ambiguidade")
    caminho = _docx(tmp_path, ["Critério de pontuação: qualidade (a definir)."], nome="amb.docx")
    r2 = editais.receber_versao(
        con,
        edital2,
        nome_arquivo=caminho.name,
        conteudo=caminho.read_bytes(),
        ator="teste",
    )
    req = editais.listar_requisitos(con, r2["versao_id"])[0]
    con.execute(
        "UPDATE edital_requisitos SET status = 'aprovado' WHERE id = ?",
        (req["id"],),
    )
    con.commit()
    with pytest.raises(ValueError, match="ambiguidade"):
        editais.aprovar_versao(con, r2["versao_id"], revisor="revisor")


def test_revisao_humana_preserva_interpretacao_original(con, tmp_path):
    _, recebido = _criar_e_receber(con, tmp_path)
    requisito = next(
        req
        for req in editais.listar_requisitos(con, recebido["versao_id"])
        if req["tipo"] == "papel"
    )
    original_descricao = requisito["descricao"]
    original_regra = requisito["regra"]
    editais.revisar_requisito(
        con,
        int(requisito["id"]),
        revisor="revisor",
        decisao="alterar",
        descricao="Papel PM/Scrum Sênior — 1 vaga (interpretação revisada)",
        regra={**original_regra, "vagas": 1, "revisado": True},
        observacao="ajuste de redação",
    )
    revisado = next(
        req
        for req in editais.listar_requisitos(con, recebido["versao_id"])
        if int(req["id"]) == int(requisito["id"])
    )
    assert revisado["status"] == "alterado"
    assert revisado["descricao_original"] == original_descricao
    assert revisado["regra_original"] == original_regra
    assert revisado["regra"]["revisado"] is True
    assert revisado["ambiguidade"] == 0


def test_obrigatorio_rejeitado_bloqueia_aprovacao(con, tmp_path):
    _, recebido = _criar_e_receber(con, tmp_path)
    requisitos = editais.listar_requisitos(con, recebido["versao_id"])
    obrigatorio = next(req for req in requisitos if req["obrigatorio"])
    editais.revisar_requisito(
        con,
        int(obrigatorio["id"]),
        revisor="revisor",
        decisao="rejeitar",
        observacao="não aceito pela área",
    )
    for req in requisitos:
        if int(req["id"]) == int(obrigatorio["id"]):
            continue
        editais.revisar_requisito(
            con,
            int(req["id"]),
            revisor="revisor",
            decisao="aprovar",
            observacao="ok",
        )
    with pytest.raises(ValueError, match="obrigatório"):
        editais.aprovar_versao(con, recebido["versao_id"], revisor="revisor")


def test_retificacao_cria_nova_versao_e_preserva_historico(con, tmp_path):
    edital_id, versao_original = _receber_e_ativar(con, tmp_path)
    antigos = editais.listar_requisitos(con, versao_original)
    assert antigos
    disputa = con.execute(
        "SELECT id, edital_versao_id FROM disputas WHERE edital_versao_id = ?",
        (versao_original,),
    ).fetchone()
    assert disputa is not None

    linhas_retificadas = list(LINHAS_EDITAL)
    linhas_retificadas[2] = "Papel: PM/Scrum Sênior — 2 vagas"
    linhas_retificadas[8] = "Experiência mínima: 7 anos em gestão de projetos"
    caminho = _docx(tmp_path, linhas_retificadas, nome="retificacao.docx")
    retificacao = editais.receber_versao(
        con,
        edital_id,
        nome_arquivo=caminho.name,
        conteudo=caminho.read_bytes(),
        ator="teste",
        tipo="retificacao",
        motivo="correção de vagas e experiência",
    )
    assert retificacao["duplicado"] is False
    assert retificacao["versao_id"] != versao_original
    assert retificacao["efeito"]["alterados"]
    assert any(
        "1 vaga" in item["antes"] and "2 vaga" in item["depois"]
        for item in retificacao["efeito"]["alterados"]
    )
    assert retificacao["efeito"]["removidos"] == []

    # A versão antiga continua com os requisitos originais.
    antigos_depois = editais.listar_requisitos(con, versao_original)
    assert [r["descricao"] for r in antigos] == [r["descricao"] for r in antigos_depois]
    versao_antiga = editais.obter_versao(con, versao_original)
    assert versao_antiga["status"] == "ativa"

    _revisar_todos(con, retificacao["versao_id"])
    editais.aprovar_versao(con, retificacao["versao_id"], revisor="revisor")
    resultado = editais.ativar_versao(con, retificacao["versao_id"], ator="teste")
    assert resultado["anterior_id"] == versao_original
    assert editais.obter_versao(con, versao_original)["status"] == "substituida"
    assert editais.obter_versao(con, retificacao["versao_id"])["status"] == "ativa"

    # A demanda anterior é preservada e passa a apontar para a nova versão,
    # com uma pendência explícita para revisão do time.
    disputa_depois = con.execute(
        "SELECT edital_versao_id FROM disputas WHERE id = ?", (disputa["id"],)
    ).fetchone()
    assert disputa_depois["edital_versao_id"] == retificacao["versao_id"]
    pendencias = con.execute(
        "SELECT * FROM atualizacoes_pendentes WHERE disputa_id = ?",
        (disputa["id"],),
    ).fetchall()
    assert pendencias
    assert pendencias[0]["candidato_id"] is None
    assert pendencias[0]["tipo"] == "edital_retificado"
    assert json.loads(pendencias[0]["detalhes_json"])["alterados"]
    listadas = listar_atualizacoes_pendentes(con, disputa_id=int(disputa["id"]))
    assert listadas
    assert listadas[0]["nome_pseudonimo"] == "Demanda (edital)"


def test_edital_com_cobertura_suficiente_cria_demanda(con, tmp_path):
    edital_id, versao_id = _receber_e_ativar(con, tmp_path)
    cobertura = editais.avaliar_cobertura(con, versao_id)
    assert cobertura["suficiente"] is True
    assert cobertura["papeis"] == 4
    assert cobertura["vagas"] == 5
    assert cobertura["empresa"] >= 1
    assert cobertura["equipe"] >= 1
    disputa = con.execute(
        "SELECT * FROM disputas WHERE edital_versao_id = ?", (versao_id,)
    ).fetchone()
    assert disputa is not None
    assert disputa["nome_ficticio"] == "Edital de teste"
    assert editais.obter_edital(con, edital_id)["status"] == "ativo"
